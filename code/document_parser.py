#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
function description: 此文件用于文档解析统一工厂（支持 txt, md, pdf, docx, csv, xlsx, xls）
author: TangKan
contact: 785455964@qq.com
IDE: PyCharm Community Edition 2026.1.1
time: 2026/06/01 15:31
version: V1.0
Target Python:  3.14 -> 3.12
"""

import os
import pandas as pd
import docx
import fitz  # PyMuPDF
import json


def parse_pdf(file_path: str) -> str:
    """解析 PDF 文件，提取所有页面的文本"""
    text_content = []
    try:
        # 打开 PDF 文件
        with fitz.open(file_path) as pdf_document:
            for page_num in range(len(pdf_document)):
                page = pdf_document.load_page(page_num)
                # 提取每一页的纯文本
                text_content.append(page.get_text("text"))
        return "\n".join(text_content)
    except Exception as e:
        return f"[PDF解析错误]: {str(e)}"


def parse_docx(file_path: str) -> str:
    """解析 Word (.docx) 文件，提取段落和表格"""
    try:
        doc = docx.Document(file_path)
        full_text = []

        # 1. 提取所有段落
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())

        # 2. 提取表格（大模型对格式化的表格数据非常敏感，如石窟监测数据表）
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))

        return "\n".join(full_text)
    except Exception as e:
        return f"[DOCX解析错误]: {str(e)}"


def parse_spreadsheet(file_path: str, ext: str) -> str:
    """解析表格文件 (CSV/Excel) 并转换为 Markdown 格式字符串供大模型阅读"""
    try:
        if ext == '.csv':
            # CSV 没有 Sheet 的概念，直接读
            df = pd.read_csv(file_path)  # type: ignore

            if not isinstance(df, pd.DataFrame):
                return "[表格解析错误]: 读取的文件内容为空或格式异常"

            return df.to_markdown(index=False)

        else:  # .xlsx 或 .xls
            # 🌟 核心修改：sheet_name=None 会一次性读取所有 Sheet
            # 返回的数据结构是: {'Sheet1': df1, 'Sheet2': df2}
            sheet_dict = pd.read_excel(file_path, sheet_name=None)  # type: ignore

            markdown_results = []

            # 遍历所有的 Sheet
            for sheet_name, df in sheet_dict.items():
                # 如果这个表是空的，或者格式不对，直接跳过
                if not isinstance(df, pd.DataFrame) or df.empty:
                    continue

                # 🌟 给大模型加上明确的层级标题，防止数据混乱
                markdown_results.append(f"### 表格子页 (Sheet): {sheet_name}")
                markdown_results.append(df.to_markdown(index=False))
                markdown_results.append("\n")  # 加上空行，视觉分割更清晰

            if not markdown_results:
                return "[表格解析错误]: Excel 文件中没有有效数据"

            # 把所有 Sheet 的 Markdown 拼成一个超长字符串
            return "\n".join(markdown_results)

    except Exception as e:
        return f"[表格解析错误]: {str(e)}"


def parse_plain_text(file_path: str) -> str:
    """解析纯文本文件"""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        # 如果 utf-8 失败，尝试 gbk (处理一些早期的中文 txt)
        with open(file_path, "r", encoding="gbk") as f:
            return f.read()
    except Exception as e:
        return f"[纯文本解析错误]: {str(e)}"


def parse_json(file_path: str) -> str:
    """解析 JSON 文件，并将其格式化为带有缩进的字符串，方便大模型理解层级关系"""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        # 🌟 核心细节：
        # indent=2 保证输出有漂亮的缩进，大模型特别吃这一套
        # ensure_ascii=False 极其关键，否则里面的中文会变成 \u4e2d 这种机器码，大模型容易糊涂
        return json.dumps(data, ensure_ascii=False, indent=2)

    except json.JSONDecodeError as e:
        return f"[JSON解析错误]: 文件格式损坏或不规范 - {str(e)}"
    except Exception as e:
        return f"[JSON读取错误]: {str(e)}"


def read_any_file(file_path: str) -> str:
    """
    🌟 统一对外接口：传入任意支持的文件路径，自动路由到对应的解析器，返回纯文本
    """
    if not os.path.exists(file_path):
        return f"⚠️ 文件不存在: {file_path}"

    # 获取文件后缀名并转为小写 (例如: '.pdf')
    ext = os.path.splitext(file_path)[1].lower()

    if ext in ['.txt', '.md']:
        return parse_plain_text(file_path)

    elif ext == '.pdf':
        return parse_pdf(file_path)

    elif ext == '.docx':
        return parse_docx(file_path)

    elif ext == '.doc':
        return "⚠️ 警告：暂不支持解析古老的 .doc 格式，请用 Word 将其另存为 .docx 后重试。"

    elif ext in ['.csv', '.xlsx', '.xls']:
        return parse_spreadsheet(file_path, ext)

    elif ext == '.json':
        return parse_json(file_path)

    else:
        return f"⚠️ 暂不支持解析此类型的文件: {ext}"


# 测试代码 (当直接运行此文件时执行)
if __name__ == "__main__":
    # 可以随便丢一个 pdf 或 xlsx 进去测试一下提取效果
    print(read_any_file("test2/parser_test/default_chat.md"))
    pass

